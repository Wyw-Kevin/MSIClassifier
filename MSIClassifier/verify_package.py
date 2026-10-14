"""Check preserved sources, package inventory, Word structure and English tables."""
from pathlib import Path
import argparse,hashlib,json,re
from zipfile import ZipFile
from lxml import etree as E
import openpyxl

ROOT=Path(__file__).resolve().parent
sha=lambda p:hashlib.sha256(p.read_bytes()).hexdigest()

def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--output',type=Path,default=ROOT/'output/package_verification.json')
    args=parser.parse_args();checks=0
    for file in ['root_source_manifest.json','reporting_extra_sources.json']:
        for r in json.loads((ROOT/'provenance'/file).read_text()):
            assert sha(ROOT/r['package'])==r['sha256'],r['package'];checks+=1
    for folder,key,path_key,hash_key in [
      ('scRNA','records','destination_relative','packaged_sha256'),
      ('spatial','copied_files','path','sha256'),
      ('bayesprism','source_records','destination','packaged_sha256')]:
        base=ROOT/'code'/folder
        provenance=base/('provenance_'+folder+'.json')
        for r in json.loads(provenance.read_text())[key]:
            assert sha(base/r[path_key])==r[hash_key],(folder,r[path_key]);checks+=1
    inventory=ROOT/'provenance/file_manifest.json'
    if inventory.exists():
        for r in json.loads(inventory.read_text())['files']:
            assert sha(ROOT/r['path'])==r['sha256'],r['path'];checks+=1
    python_files=[]
    for p in ROOT.rglob('*.py'):
        if '__pycache__' in p.parts:continue
        compile(p.read_bytes(),str(p),'exec');python_files.append(str(p.relative_to(ROOT)))
    W='http://schemas.openxmlformats.org/wordprocessingml/2006/main';ns={'w':W}
    with ZipFile(ROOT/'supplement/MSIClassifier_补充方法与图片.docx') as z:
        assert z.testzip() is None
        r=E.fromstring(z.read('word/document.xml'))
        assert len(r.findall('.//{'+W+'}drawing'))==3
        assert not r.findall('.//{'+W+'}tbl')
        for para in r.findall('.//{'+W+'}p'):
            if not ''.join(para.xpath('.//w:t/text()',namespaces=ns)):continue
            s=para.find('w:pPr/w:spacing',ns)
            assert s.get('{'+W+'}line')=='480' and s.get('{'+W+'}lineRule')=='exact'
    book=openpyxl.load_workbook(ROOT/'supplement/SupplementaryTablesS1-S14_English.xlsx',read_only=True)
    assert len(book.sheetnames)==37
    cells=0
    for sheet in book:
        for row in sheet.values:
            for v in row:
                if isinstance(v,str):assert not re.search(r'[\u4e00-\u9fff]',v),(sheet.title,v[:50])
                cells+=1
    result={'status':'passed','source_and_inventory_hash_checks':checks,'python_sources_compiled':len(python_files),'supplement_figures':3,'supplement_tables_in_word':0,'table_workbook_sheets':len(book.sheetnames),'table_cells_checked':cells,'tables_all_English':True,'supplement_body_line_spacing':'fixed24pt','scope':'Package integrity and stored document structure; independent of raw biological validation'}
    args.output.parent.mkdir(parents=True,exist_ok=True);args.output.write_text(json.dumps(result,indent=2)+'\n')
    print(json.dumps(result,indent=2))

if __name__=='__main__':main()
