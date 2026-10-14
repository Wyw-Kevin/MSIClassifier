"""Publish verified scRNA sensitivity results into a new manuscript and supplement."""
from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED
from copy import deepcopy
import json, hashlib, shutil, re
import pandas as pd
import numpy as np
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from lxml import etree as E
from PIL import Image

HERE = Path(__file__).resolve().parent
BASE = HERE.parent/'第三版_十亚型与留一队列MMR4修订_研究设计图恢复版'
PREVIOUS = BASE/'补充图S2十亚型数量入选率'
SOURCE = PREVIOUS/'MMR分模态分类器_第三版_补充图S2十亚型数量入选率_句子标红.docx'
DEST = HERE/'MMR分模态分类器_第三版_AUROC阈值全候选敏感性分析_句子标红.docx'
SUP = HERE/'supplement'
FIG = HERE/'figures'
sha = lambda path: hashlib.sha256(path.read_bytes()).hexdigest()
verification = json.loads((HERE/'verification.json').read_text())
assert verification['status'] == 'passed' and verification['fitted_models_replayed'] == 42
assert sha(SOURCE) == '598472145d8620d80d1b497401df273b5a83d821fc7fe14344ef6056c1d8bfb9'
SUP.mkdir(exist_ok=True)
FIG.mkdir(exist_ok=True)
for path in (PREVIOUS/'supplement').iterdir():
    if path.is_file():
        shutil.copy2(path, SUP/path.name)

# Supplemental table extensions preserve the original measurement tables.
threshold_table = pd.read_csv(HERE/'tables/S10_threshold_performance.csv')
screen = threshold_table.loc[threshold_table.family.eq('subtype_threshold')].copy()
fixed18 = threshold_table.loc[threshold_table.family.eq('fixed_panel_reference')].copy()
cnv = pd.read_csv(HERE/'tables/S3_CNV_quantile_performance.csv')
ranges = pd.read_csv(HERE/'tables/AUROC_sensitivity_ranges.csv')
tables = {
    'S10_thresholds': ('S10_threshold_performance.csv', threshold_table),
    'S10_threshold_membership': ('S10_threshold_membership.csv', pd.read_csv(HERE/'tables/subtype_threshold_membership.csv')),
    'S2_CNV_AUROC_ranges': ('S2_CNV_AUROC_ranges.csv', ranges.loc[ranges.family.eq('CNV_quantile')]),
    'S3_CNV_AUROC': ('S3_CNV_quantile_performance.csv', cnv),
}
for filename, frame in tables.values():
    frame.to_csv(SUP/filename, index=False)
screen_ranges = ranges.loc[ranges.family.eq('subtype_threshold')]
screen_ranges.to_csv(SUP/'S10_threshold_AUROC_ranges.csv', index=False)
tables['S10_threshold_ranges'] = ('S10_threshold_AUROC_ranges.csv', screen_ranges)
book_path = SUP/'SupplementaryTablesS1-S14_English.xlsx'
book = openpyxl.load_workbook(book_path)
old_values = {s.title: list(s.values) for s in book}
descriptions = {
    'S10_thresholds': 'scRNA cutoff sensitivity: all passing candidates, no count cap; historical top18 reference separately labeled',
    'S10_threshold_membership': '71 reference candidates at oriented AUROC cutoffs 0.60/0.65/0.70/0.75',
    'S2_CNV_AUROC_ranges': 'CNV q90/q95/q99 AUROC ranges: refitted and frozen-q95 classifiers',
    'S3_CNV_AUROC': 'CNV quantile per-cohort performance: refitted and frozen-q95 classifiers',
    'S10_threshold_ranges': 'scRNA cutoff per-cohort AUROC ranges; baseline is24-subtype cutoff0.65, distinct from historical top18',
}
for name, (filename, frame) in tables.items():
    ws = book.create_sheet(name)
    ws.append(frame.columns.tolist())
    for row in frame.replace({np.nan:None}).itertuples(index=False, name=None):
        ws.append(list(row))
    for cell in ws[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='38566B')
    for column in ws.columns:
        ws.column_dimensions[get_column_letter(column[0].column)].width = min(45, max(14, len(str(column[0].value))+2))
    for row in ws:
        for cell in row:
            cell.alignment = Alignment(vertical='top', wrap_text=True)
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions
    book['Readme'].append([name, descriptions[name]])
book.save(book_path)

# Figure S3 reports only refitted comparisons; frozen-q95 results are tabulated.
plt.rcParams.update({'font.family':'DejaVu Sans','font.size':10,
    'svg.fonttype':'none','pdf.fonttype':42,'axes.spines.top':False,'axes.spines.right':False})
fig, axes = plt.subplots(1, 2, figsize=(12, 4.8), sharey=True)
colors = {'GSE178341':'#226696','GSE132465':'#b56327','GSE200997':'#39889c'}
labels = {'GSE178341':'GSE178341 grouped OOF (n=64)',
          'GSE132465':'GSE132465 retrospective (n=23)',
          'GSE200997':'GSE200997 retrospective (n=15)'}
for ds, color in colors.items():
    a = screen.loc[screen.dataset.eq(ds)].sort_values('screening_cutoff')
    b = cnv.loc[cnv.dataset.eq(ds)&cnv['mode'].eq('refit')].sort_values('CNV_quantile')
    axes[0].plot(a.screening_cutoff, a.AUROC, color=color, marker='o', label=labels[ds], lw=1.8)
    baseline = float(fixed18.loc[fixed18.dataset.eq(ds),'AUROC'].iloc[0])
    axes[0].scatter([.65],[baseline],color=color,marker='*',s=180,edgecolors='white',linewidths=.6,zorder=5)
    axes[1].plot(b.CNV_quantile, b.AUROC, color=color, marker='o', lw=1.8)
axes[0].axvline(.65, color='#888888', ls='--', lw=1)
axes[1].axvline(.95, color='#888888', ls='--', lw=1)
axes[0].set(xlabel='Oriented AUROC screening cutoff', ylabel='AUROC',
            xticks=[.60,.65,.70,.75], ylim=(.89,1.01))
axes[1].set(xlabel='Stromal-reference CNV quantile', xticks=[.90,.95,.99])
axes[1].set_xticklabels(['q90','q95','q99'])
axes[0].set_title('A   Subtype-screening threshold', loc='left', weight='bold', pad=30)
axes[1].set_title('B   CNV candidate definition', loc='left', weight='bold', pad=30)
axes[0].text(.5,1.035,'All passing candidates: 39 / 24 / 17 / 9 (no count cap)',
             transform=axes[0].transAxes,ha='center',fontsize=8.5)
axes[1].text(.5,1.035,'Fixed 18 subtypes; weighted MLH1 recalculated at each quantile',
             transform=axes[1].transAxes,ha='center',fontsize=8.5)
fig.legend(*axes[0].get_legend_handles_labels(), loc='upper center',
           bbox_to_anchor=(.5,.975), ncol=3, frameon=False, fontsize=9)
fig.subplots_adjust(left=.065,right=.98,top=.78,bottom=.25,wspace=.16)
fig.text(.065,.085,
    'Archived patient-grouped folds; training-only Z-score and shrinkage LDA. Historical candidate ranking and cohort-level CNV fixed.\n'
    'Stars: historical top 18 reference at 0.65; circles: all passing candidates. Dashed lines: 0.65/q95. GSE200997: two positive cases.\n'
    'Frozen-q95 input sensitivity is reported in Table S3.',
    fontsize=8.5)
for ext in ['png','svg','pdf']:
    filename = 'SupplementaryFigureS3_scRNA_threshold_CNV_sensitivity.'+ext
    fig.savefig(FIG/filename, dpi=260, bbox_inches='tight')
    shutil.copy2(FIG/filename, SUP/filename)
plt.close(fig)

NS = {'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main',
      'r':'http://schemas.openxmlformats.org/officeDocument/2006/relationships',
      'wp':'http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing',
      'a':'http://schemas.openxmlformats.org/drawingml/2006/main',
      'pic':'http://schemas.openxmlformats.org/drawingml/2006/picture',
      'm':'http://schemas.openxmlformats.org/officeDocument/2006/math'}
q = lambda name:'{'+NS['w']+'}'+name
text = lambda p:''.join(p.xpath('.//w:t/text()',namespaces=NS))
with ZipFile(SOURCE) as archive:
    infos = archive.infolist()
    original = {info.filename:archive.read(info.filename) for info in infos}
root = E.fromstring(original['word/document.xml'])
body = root.find('w:body',NS)
blocks = list(body)
replacements = {
60:[('q90、q95、q99及边界规则用于测量层面的敏感性分析，考察候选细胞数和特征值变化。',
     'q90、q95、q99及边界规则的候选细胞计数和特征值比较见补充表S2–S3，相应分类性能敏感性分析见补充图S3。')],
135:[('q90/q95/q99敏感性分析目前限于测量层面。',
      'q90/q95/q99候选定义变化对AUROC的影响存在队列差异，GSE132465在q99下的判别性能降低。敏感性分析以既定注释、CNV分数和患者折分为条件，上游测量的不确定性仍未纳入。')],
140:[('S1–S5记录样本QC、CNV、表达层与标签来源；S6–S9列示配对差值、特征对照、样本纳入与低细胞数敏感性；S10–S11提供候选筛选及单亚型方向记录；S12–S13记录MMR4留一队列阈值及11种分类器的逐队列结果；S14列示全部候选的2,000次bootstrap入选频率、入选数量、阈值敏感性、集合重合度及原十亚型集合的完全复现情况。',
      'S1–S5记录样本QC、CNV、表达层与标签来源，S2–S3附表列示CNV阈值对应的AUROC及变化范围。S6–S9列示配对差值、特征对照、样本纳入与低细胞数敏感性。S10–S11提供候选筛选、单亚型方向及单细胞筛选阈值敏感性结果。S12–S13记录MMR4留一队列阈值及11种分类器的逐队列结果。S14列示全部候选的2,000次bootstrap入选频率、入选数量、阈值敏感性、集合重合度及原十亚型集合的完全复现情况。')]
}
# Explain candidate screening and historical panel-size choice as distinct steps.
for index, block in enumerate(blocks):
    original_text = text(block)
    first = '以GSE178341方向校正AUROC>0.65为探索性筛选标准保留24个亚型，再按该AUROC降序取前18个。'
    second = '18亚型面板源于既往10–24亚型配置比较，其开发过程涉及GSE132465结果。'
    if first in original_text:
        replacements.setdefault(index, []).extend([
            (first, '以GSE178341方向校正AUROC>0.65为探索性筛选标准获得24个候选亚型，并按该指标降序构建包含前10至24个亚型的嵌套特征集。'),
            (second, '在既往联合全上皮MLH1的分类器配置比较中保留18亚型收缩LDA方案，其亚型组成沿用于当前加权MLH1模型；该探索阶段同时考察了GSE178341和GSE132465的表现。')
        ])
        break
else:
    raise AssertionError('Historical panel-selection paragraph not found')

audit, ranges_equal, ranges_red = [], {}, {}
def red_run(p, value, template):
    rp = deepcopy(template)
    for color in rp.findall('w:color',NS):
        rp.remove(color)
    E.SubElement(rp,q('color'),{q('val'):'FF0000'})
    for sentence in re.findall(r'[^。！？]+[。！？]?|[。！？]',value):
        run = E.SubElement(p,q('r'))
        run.append(deepcopy(rp))
        t = E.SubElement(run,q('t'))
        t.set('{http://www.w3.org/XML/1998/namespace}space','preserve')
        t.text = sentence

def revise(p, changes, index):
    old = text(p)
    runs, offset = [], 0
    for run in p.findall('w:r',NS):
        ts = run.findall('w:t',NS)
        assert len(ts)==1 and set(E.QName(c).localname for c in run)<={'rPr','t'}
        value = ts[0].text or ''
        runs.append((offset,offset+len(value),run))
        offset += len(value)
    assert offset==len(old)
    spans=[]
    for previous,revised in changes:
        assert old.count(previous)==1
        start=old.index(previous)
        spans.append((start,start+len(previous),previous,revised))
    spans.sort()
    for run in p.findall('w:r',NS):
        p.remove(run)
    cursor, position, equal, red = 0,0,[],[]
    def keep(start,end):
        nonlocal position
        new_start=position
        for a,b,run in runs:
            lo,hi=max(a,start),min(b,end)
            if lo<hi:
                clone=deepcopy(run)
                clone.find('w:t',NS).text=old[lo:hi]
                clone.find('w:t',NS).set('{http://www.w3.org/XML/1998/namespace}space','preserve')
                p.append(clone)
        position += end-start
        equal.append((start,end,new_start,position))
    for start,end,previous,revised in spans:
        keep(cursor,start)
        template=next(run.find('w:rPr',NS) for a,b,run in runs if a<=start<b)
        red_start=position
        red_run(p,revised,template)
        position += len(revised)
        red.append((red_start,position))
        cursor=end
        audit.append({'original_block':index,'before':previous,'after':revised})
    keep(cursor,len(old))
    ranges_equal[index],ranges_red[index]=equal,red

for index,changes in replacements.items():
    revise(blocks[index],changes,index)

def new_paragraph(value, template, identity):
    p=E.Element(q('p'))
    p.append(deepcopy(template.find('w:pPr',NS)))
    rp=template.find('w:r/w:rPr',NS)
    assert rp is not None
    red_run(p,value,rp)
    audit.append({'new_paragraph':identity,'after':value})
    return p

method_threshold=('单细胞亚型筛选阈值敏感性分析固定GSE178341历史方向及定向AUROC排序，分别采用0.60、0.65、0.70和0.75界值。'
    '各界值下纳入全部达标亚型，分别形成39、24、17和9亚型面板，均联合q95加权MLH1；原18亚型模型作为固定面板参照单独列示。'
    '各面板沿用原患者五折划分，标准化及收缩LDA仅在相应训练折拟合，完整开发集模型用于两个回顾性队列评分。'
    '候选方向及AUROC取自完整开发队列，上游注释与CNV测量保持固定，所得结果用于描述AUROC筛选阈值的条件性敏感性。')
method_cnv=('CNV阈值敏感性分析固定18亚型，分别按各队列基质参考分数的q90、q95和q99界定候选恶性上皮细胞。'
    '各阈值下的加权MLH1均由同一CNV保留集合的表达及细胞计数计算。'
    '各方案沿用相同患者折分，分别在训练折拟合标准化和收缩LDA，完整开发集模型用于回顾性队列评分。'
    '另将q90和q99输入送入冻结的q95折内及完整开发集分类器，评价候选定义改变造成的输入扰动。'
    '主分析沿用既往配置比较后保留的18亚型面板与q95候选定义。')
result_threshold=('AUROC筛选阈值为0.60、0.65、0.70和0.75时，全部达标亚型数分别为39、24、17和9。'
    '相应联合模型在GSE178341患者分组OOF中的AUROC范围为0.9350–0.9803，在GSE132465中为0.9474–0.9737，在GSE200997中为0.9231–1.0000（补充表S10、补充图S3A）。'
    '0.65阈值下的24亚型模型在三个队列中的AUROC分别为0.9547、0.9605和0.9615；作为固定面板参照的原18亚型模型分别为0.9783、0.9605和1.0000。'
    '0.70阈值下的17亚型模型对应AUROC分别为0.9803、0.9474和1.0000，提高筛选阈值未带来跨队列一致的AUROC增益。')
result_cnv=('q90、q95和q99规则下，三个单细胞队列均分别有64、23和15例可评分。'
    'q90与q95对应联合模型在GSE178341 OOF、GSE132465和GSE200997中的AUROC均分别为0.9783、0.9605和1.0000。'
    'q99对应AUROC分别为0.9793、0.9079和1.0000，较q95分别变化0.0010、−0.0526和0。'
    '三个队列的CNV阈值敏感性AUROC范围分别为0.9783–0.9793、0.9079–0.9605和1.0000（补充表S2–S3、补充图S3B）。'
    '冻结q95分类器的输入扰动分析得到相同的逐队列AUROC，样本评分另列于重现资料。'
    'GSE132465在q99下的AUROC降低，显示候选恶性上皮定义的影响存在队列差异。')
new_body = []
for index,values in [(60,[method_threshold,method_cnv]),(79,[result_threshold]),(88,[result_cnv])]:
    anchor=blocks[index]
    for ordinal,value in enumerate(values):
        p=new_paragraph(value,blocks[index],f'after{index}:{ordinal}')
        body.insert(body.index(anchor)+1,p)
        anchor=p
        new_body.append(p)
caption=('补充图S3. 单细胞亚型筛选与CNV阈值的条件性敏感性分析。'
    '（A）以0.60、0.65、0.70和0.75为筛选界值，纳入全部达标亚型，分别为39、24、17和9个，并联合q95加权MLH1。圆点表示各阈值下的全部达标亚型模型，星号单列原18亚型模型作为固定面板参照。'
    '（B）固定18亚型，分别使用q90、q95和q99定义的加权MLH1重新拟合分类器。'
    '各方案使用同一患者五折，标准化和收缩LDA仅在训练折拟合；回顾性队列使用完整开发集模型。'
    '虚线标示0.65筛选界值和q95候选定义。'
    'GSE178341、GSE132465和GSE200997分别有64、23和15例，GSE200997仅2例阳性。'
    '冻结q95分类器的输入扰动结果见补充表S3。')
caption_p=new_paragraph(caption,blocks[143],'FigureS3_caption')
body.insert(body.index(blocks[144])+1,caption_p)
drawing=deepcopy(blocks[144])
new_rid='rIdScRNASensitivityS3PNG'
for node in drawing.xpath('.//a:blip',namespaces=NS):
    node.set('{'+NS['r']+'}embed',new_rid)
for node in drawing.xpath('.//wp:docPr',namespaces=NS):
    max_id=max(int(e.get('id')) for e in root.xpath('//wp:docPr',namespaces=NS))
    node.set('id',str(max_id+1))
    node.set('name','Supplementary Figure S3')
for node in drawing.xpath('.//pic:cNvPr',namespaces=NS):
    node.set('name','Supplementary Figure S3')
png=FIG/'SupplementaryFigureS3_scRNA_threshold_CNV_sensitivity.png'
width,height=Image.open(png).size
extent=drawing.find('.//wp:extent',NS)
cx=int(extent.get('cx'))
cy=round(cx*height/width)
extent.set('cy',str(cy))
for ext in drawing.xpath('.//a:xfrm/a:ext',namespaces=NS):
    ext.set('cx',str(cx));ext.set('cy',str(cy))
body.insert(body.index(caption_p)+1,drawing)
new_blocks=list(body)
old_to_new={i:new_blocks.index(p) for i,p in enumerate(blocks)}
rels=E.fromstring(original['word/_rels/document.xml.rels'])
package_ns='http://schemas.openxmlformats.org/package/2006/relationships'
target='media/scRNA_sensitivity_S3.png'
E.SubElement(rels,'{'+package_ns+'}Relationship',{
    'Id':new_rid,'Type':NS['r']+'/image','Target':target})
updated=dict(original)
updated['word/document.xml']=E.tostring(root,xml_declaration=True,encoding='UTF-8',standalone=True)
updated['word/_rels/document.xml.rels']=E.tostring(rels,xml_declaration=True,encoding='UTF-8',standalone=True)
updated['word/'+target]=png.read_bytes()
with ZipFile(DEST,'w') as archive:
    for info in infos:
        archive.writestr(info,updated[info.filename])
    archive.writestr('word/'+target,updated['word/'+target])

# Verify saved DOCX, all pre-existing table values, media, equations and format.
with ZipFile(DEST) as archive:
    assert archive.testzip() is None
    saved={name:archive.read(name) for name in archive.namelist()}
assert set(saved)==set(original)|{'word/'+target}
assert all(saved[name]==value for name,value in original.items()
           if name not in ['word/document.xml','word/_rels/document.xml.rels'])
old_root=E.fromstring(original['word/document.xml'])
saved_root=E.fromstring(saved['word/document.xml'])
old_blocks=list(old_root.find('w:body',NS))
saved_blocks=list(saved_root.find('w:body',NS))
assert len(saved_blocks)==len(old_blocks)+6
for i,old in enumerate(old_blocks):
    new=saved_blocks[old_to_new[i]]
    if i not in replacements:
        assert E.tostring(old,method='c14n')==E.tostring(new,method='c14n')
    else:
        assert E.tostring(old.find('w:pPr',NS))==E.tostring(new.find('w:pPr',NS))
        expected=text(old)
        for previous,revised in replacements[i]:
            expected=expected.replace(previous,revised)
        assert text(new)==expected

def char_styles(p):
    values=[]
    for run in p.findall('w:r',NS):
        rp=run.find('w:rPr',NS)
        style=E.tostring(rp,method='c14n')
        color=rp.find('w:color',NS)
        values.extend((ch,style,color.get(q('val')) if color is not None else None) for ch in text(run))
    return values
for i in replacements:
    old,new=char_styles(old_blocks[i]),char_styles(saved_blocks[old_to_new[i]])
    for a,b,c,d in ranges_equal[i]:
        assert old[a:b]==new[c:d]
    for start,end in ranges_red[i]:
        assert all(color=='FF0000' for _,_,color in new[start:end])
for p in new_body+[caption_p]:
    p=saved_blocks[new_blocks.index(p)]
    assert all(c=='FF0000' for _,_,c in char_styles(p))
    spacing=p.find('w:pPr/w:spacing',NS)
    assert spacing.get(q('line'))=='480' and spacing.get(q('lineRule'))=='exact'
    for size in p.xpath('.//w:rPr/w:sz',namespaces=NS):
        assert size.get(q('val'))=='24'
assert len(saved_root.xpath('//m:oMath',namespaces=NS))==5
assert len(saved_root.xpath('//w:tbl',namespaces=NS))==6
for p in saved_root.xpath('//w:tbl//w:p',namespaces=NS):
    spacing=p.find('w:pPr/w:spacing',NS)
    assert spacing.get(q('line'))=='240' and spacing.get(q('lineRule'))=='auto'
assert saved['word/'+target]==png.read_bytes()
new_book=openpyxl.load_workbook(book_path)
for name,values in old_values.items():
    rows=list(new_book[name].values)
    assert rows[:len(values)]==values
    if name!='Readme':
        assert rows==values
for name,(_,frame) in tables.items():
    rows=list(new_book[name].values)
    assert rows[0]==tuple(frame.columns)
    for old,new in zip(frame.itertuples(index=False,name=None),rows[1:]):
        for a,b in zip(old,new):
            if isinstance(a,(float,np.floating)):
                assert abs(a-b)<1e-12
            else:
                assert a==b
for sheet in new_book:
    for row in sheet:
        for cell in row:
            assert not(isinstance(cell.value,str) and re.search(r'[\u3400-\u9fff]',cell.value))
assert sha(SOURCE)=='598472145d8620d80d1b497401df273b5a83d821fc7fe14344ef6056c1d8bfb9'
(HERE/'revision_audit.json').write_text(json.dumps(audit,ensure_ascii=False,indent=2))
(HERE/'delivery_verification.json').write_text(json.dumps({
    'status':'passed','source':str(SOURCE),'output':str(DEST),'output_SHA256':sha(DEST),
    'new_body_paragraphs':4,'new_supplementary_figure':'S3',
    'modified_original_blocks':sorted(replacements),'all_original_images_tables_equations':'unchanged',
    'new_sentences_red':'verified','body_and_captions':'12pt, fixed24pt; body first-line indent preserved',
    'table_line_spacing':'single','all_existing_workbook_values':'preserved',
    'new_English_supplement_sheets':list(tables),'analysis_verification_checks':verification['checks']
},ensure_ascii=False,indent=2))
(HERE/'README.txt').write_text(
    'Completed exploratory scRNA sensitivity analysis, seed2026. Main classifier and its primary q95/0.65 identity preserved.\n'
    'Subtype cutoffs0.60/0.65/0.70/0.75 use fixed historical full-development ranking, strict cutoff and all passing candidates39/24/17/9; no count cap. Historical top18/q95 is a separately labeled reference.\n'
    'CNV q90/q95/q99 analyses keep all18 subtypes. Refit comparisons fit training-only scalers/LDA in archived patient folds; frozen-q95 comparisons change only measurement inputs.\n'
    'External cohorts are scored without using labels to choose panels, quantiles or preprocessing. Historically viewed cohorts remain retrospective.\n'
    '42 new fitted artifacts and1660 saved predictions passed independent saved-model, pairwise-AUROC and scaler/patient audits. Main-rule predictions reproduce archived scores within1e-10.\n'
    'The original S2/S3 measurement tables and all previous supplemental tables are preserved. New S2/S3 performance extensions and S10 threshold sheets are in the English workbook.\n'
    'Run run_analysis.py, then verify_analysis.py to reproduce numeric results; source_archive contains the required small inputs and original primary classifier. Delivery build uses the prior manuscript/supplement in the project.\n')
readme=(SUP/'README.txt').read_text()
(SUP/'README.txt').write_text(readme+'\nNew scRNA sensitivity extensions: S10_thresholds/membership/ranges; S2_CNV_AUROC_ranges; S3_CNV_AUROC. Figure S3 shows the refitted comparisons. Historical top18 reference, CNV analyses and all main performance identities remain unchanged. Threshold-only0.65 refers to24 subtypes.\n')
archive_path=HERE/'AUROC阈值全候选敏感性分析_稿件及补充资料.zip'
with ZipFile(archive_path,'w',ZIP_DEFLATED) as archive:
    for directory in ['inputs','models','tables','supplement','figures']:
        for path in sorted((HERE/directory).rglob('*')):
            if path.is_file():
                archive.write(path,path.relative_to(HERE))
    for path in [DEST,*HERE.glob('*.py'),HERE/'README.txt',HERE/'protocol.json',HERE/'source_manifest.json',
                 HERE/'verification.json',HERE/'transfer_coverage_audit.json',HERE/'portable_replay_verification.json',HERE/'prior_comparison_verification.json',HERE/'regression_evidence.json',HERE/'delivery_verification.json',HERE/'revision_audit.json',HERE/'model_index.csv',HERE/'all_predictions.csv']:
        archive.write(path,path.name)
with ZipFile(archive_path) as archive:
    assert archive.testzip() is None
print(json.dumps({'status':'passed','manuscript':str(DEST),'supplement':str(book_path),
    'figure':str(FIG/'SupplementaryFigureS3_scRNA_threshold_CNV_sensitivity.png'),
    'archive':str(archive_path),'new_English_sheets':list(tables)},ensure_ascii=False,indent=2))
