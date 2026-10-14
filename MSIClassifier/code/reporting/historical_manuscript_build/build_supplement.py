"""Extend S12/S13 and provide reproducible implementation details in English."""
from pathlib import Path
import json, re, shutil
import pandas as pd
import openpyxl
from zipfile import ZipFile, ZIP_DEFLATED
from lxml import etree as E
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

HERE = Path(__file__).resolve().parent
OLD = HERE.parent / 'bulk12_loco_manuscript_seed2026_20261013/supplement'
OUT = HERE / 'supplement'
OUT.mkdir(exist_ok=True)
BOOK_NAME = 'SupplementaryTablesS1-S14_English.xlsx'
def prose_spacing(value):
    if not isinstance(value, str): return value
    keep = {'MMR4', 'MLH1', 'MSH2', 'MSH6', 'PMS2', 'log2', 'log1p', 'int64', 'float64', 'cell2location', 'CP10K'}
    def fix(m):
        word = m.group()
        if word in keep or re.fullmatch(r'(?:GSE|GSM|S)\d+', word): return word
        return re.sub(r'^([A-Za-z]+)(\d)', r'\1 \2', word)
    return re.sub(r'[A-Za-z]+\d+[A-Za-z]*', fix, value)
shutil.copytree(OLD, OUT, dirs_exist_ok=True, ignore=shutil.ignore_patterns('historical_bulk10', '__pycache__'))
book = openpyxl.load_workbook(OLD / BOOK_NAME)

def add_sheet(name, frame):
    if name in book.sheetnames: del book[name]
    w = book.create_sheet(name)
    w.append(list(frame.columns))
    for row in frame.where(pd.notna(frame), None).itertuples(index=False, name=None): w.append(list(row))
    w.freeze_panes = 'A2'; w.auto_filter.ref = w.dimensions
    for c in w[1]: c.font = Font(bold=True, color='FFFFFF'); c.fill = PatternFill('solid', fgColor='38566B')
    for col in w.columns: w.column_dimensions[get_column_letter(col[0].column)].width = min(54, max(15, len(str(col[0].value)) + 2))
    for row in w:
        for c in row:
            c.alignment = Alignment(vertical='top', wrap_text=True)
            if isinstance(c.value, float): c.number_format = '0.000000'

for name, file in [('S13_Fusion_vs_MMR4', 'S13_fusion_vs_MMR4_cohort.csv'),
                   ('S13_Fusion_macro', 'S13_fusion_vs_MMR4_macro.csv')]:
    frame = pd.read_csv(HERE / 'paired_increment' / file, float_precision='round_trip')
    add_sheet(name, frame)
    shutil.copy2(HERE / 'paired_increment' / file, OUT / file)
scale = pd.read_csv(HERE / 'audit_mmr4/MMR4_expression_scale_audit.csv', float_precision='round_trip')
add_sheet('S12_expression_scale', scale)
scale.to_csv(OUT / 'S12_expression_scale.csv', index=False, float_format='%.17g')

details = [
    ['BayesPrism', 'Python implementation', 'pybayesprism 0.1.0; archived local patches', 'Recorded package METADATA and preserved runtime code; not the current repository HEAD'],
    ['BayesPrism', 'Repository', 'https://github.com/ziluwang829/pyBayesPrism', 'Software author Zilu Wang; method paper reference32 and software reference37'],
    ['BayesPrism', 'Reference expression', "GSE178341.h5ad layers['counts']; raw counts", '71 non-epithelial labels plus Epi; no per-cell total normalization or logarithm'],
    ['BayesPrism', 'Reference genes', '17039 after filtering', 'Repeated symbols summed; >=5 expressing reference cells; excluded gene categories and protein-coding filter'],
    ['BayesPrism', 'Input rounding', 'numpy.rint(linear_matrix).astype(int64)', 'Whole target input rounded before Prism.new; nearest even at half integers; affects initial inference, reference update and final inference'],
    ['BayesPrism', 'Continuous target input', '*_linear.pkl', 'Preserved for audit; actual run reads *_sampling.pkl'],
    ['BayesPrism', 'Final abundance', 'posterior_theta_f.theta / final_fractions.csv', 'Updated-reference posterior mean; subtype divided by sum of same clTopLevel parent'],
    ['BayesPrism', 'Model branch', 'key=None; update_gibbs=True', 'Generic mixture branch; all reference types updated within each target cohort'],
    ['BayesPrism', 'Chain settings', 'chain_length=1000; burn=500; thin=2; alpha=1; MAP sigma=2', '250 retained states (500,502,...,998); saved implementation performs999 updates (indices0..998)'],
    ['BayesPrism', 'Historical metadata limits', 'Installation Git commit and exact dependency version snapshot not recorded', 'Preserved code hashes and patch script identify actual implementation; later environment versions do not substitute for missing historical metadata'],
    ['MMR4', 'Pre-Z-score expression', 'TCGA log2(RSEM+1); seven GEO cohorts processed log2 microarray expression', 'All4848 gene values replayed against stored expression.csv; BayesPrism inverse transform and rounding not reused'],
    ['MMR4', 'Tie rule', 'Descending rounded12 macro AUROC; descending rounded12 macro BA; ascending cutoff vector', 'Lexicographically smallest MLH1/MSH2/MSH6/PMS2 vector; rounding occurs before comparison'],
    ['scRNA', 'AUROC unit', '64 specimen scores, not62 patient averages', 'Patient grouped CV and patient cluster bootstrap; repeat specimens travel together'],
    ['scRNA', 'CNV tumor draw', 'All cells in QC-passing tumor-tissue samples', 'Not a malignant-only draw; reference-cell union performed afterward'],
    ['scRNA', 'Additional MMR gene increment', 'No archived direct comparison for current weighted MSI/dMMR classifier located', 'Four-gene expression distributions and a separate stored-MLH1-loss target do not establish this increment'],
    ['Spatial', 'Posterior export', 'export_posterior;200 draws; post_sample_means', 'Parameter means inserted into reconstruction formula; not per-draw reconstructions averaged afterward'],
    ['Spatial', 'Abundance summary', 'means_cell_abundance_w_sf', 'Posterior mean, not q05'],
    ['Spatial', 'Reference signature', 'means_per_cluster_mu_fg', 'Raw-count negative binomial reference;73 fixed types'],
    ['Spatial', 'Expression reconstruction', 'Conditional expected type-specific allocation of observed spot counts', 'Uses w_sf, m_g, s_g_gene_add and fixed reference signature'],
    ['Spatial', 'Normalization denominator', 'All12079 shared, reference-filtered model genes', 'Candidate allocated counts normalized to10000 per spot, then natural log1p'],
    ['Spatial', 'Gene pool and filter', 'Reference intersect first spatial section intersect fixed bulk-common genes;16005 unique symbols before filter', 'cell_count_cutoff=15;cell_percentage_cutoff2=0.03;nonz_mean_cutoff=1.12;12079 retained genes'],
    ['Spatial', 'Training history', 'Reference180 epochs: history_ empty; all14 section500-epoch histories preserved', 'Reference convergence not established; finite section loss traces do not by themselves prove convergence'],
    ['Paired bulk increment', 'Contrast', 'AUROC(fixed0.5 fusion)-AUROC(cohort-held-out MMR4)', 'Same sample IDs/outcomes and MMR4 held-out round; no subtype or threshold refit'],
    ['Paired bulk increment', 'Bootstrap', '2000 outcome-stratified sample resamples within cohort; RNG2026 advancing stream', 'Archived common indices reproduced; percentile95; GEO patient clustering not independently audited'],
    ['Paired bulk increment', 'Macro interval', 'Equal cohort weights; same replicate averages within-cohort differences', 'No cohort resampling; fixed observed cohorts, upstream features, predictions, scalers and cutoff locks'],
    ['Paired bulk increment', 'Display rounding', 'Point differences and intervals computed from unrounded original values', 'Retrospective delta0.03347756 rounds0.0335 although displayed0.9387-0.9053=0.0334'],
]
methods = pd.DataFrame([[prose_spacing(v) for v in row] for row in details], columns=['branch', 'item', 'actual_implementation', 'scope_or_evidence'])
add_sheet('Methods_provenance', methods)
methods.to_csv(OUT / 'Methods_provenance.csv', index=False)
coverage = pd.DataFrame([
    ['GSE24550', 13616, 13616], ['GSE25071', 9505, 9505], ['GSE27544', 14348, 14348],
    ['GSE35896', 15082, 15082], ['GSE39084', 15082, 15082], ['GSE39582', 15082, 15082],
    ['GSE75316', 15319, 15319], ['TCGA-COADREAD', 14795, 14794],
], columns=['dataset', 'initial_reference_overlap', 'actual_aligned_run_genes'])
add_sheet('BayesPrism_gene_coverage', coverage)
coverage.to_csv(OUT / 'BayesPrism_gene_coverage.csv', index=False)
for row in [
    ['S12_expression_scale', 'Pre-Z-score MMR4 inputs; original logarithmic matrices and scale replay'],
    ['S13_Fusion_vs_MMR4', 'Paired fusion-minus-MMR4 cohort AUROC differences and95% intervals'],
    ['S13_Fusion_macro', 'Equal-cohort paired macro differences and95% intervals'],
    ['Methods_provenance', 'Actual software implementation, expression layers, reconstruction and evaluation units'],
    ['BayesPrism_gene_coverage', 'Initial reference intersections and actual genes after implementation filtering'],
]: book['Readme'].append(row)
for w in book:
    for row in w:
        for c in row:
            assert not isinstance(c.value, str) or re.search(r'[\u4e00-\u9fff]', c.value) is None, (w.title, c.coordinate, c.value)
book.save(OUT / BOOK_NAME)

sections = [
    ('Python BayesPrism implementation',
     'The current twelve-subtype matrices originate from the preserved pybayesprism0.1.0 Python port (ziluwang829/pyBayesPrism), with archived local correctness and runtime patches. The method is described by Chu et al. (reference32); the software is reference37. The exact installation Git commit and historical dependency-version snapshot were not retained. Current installed versions are not substituted for those missing records. Original and modified source files, patch instructions, hashes and parameter records are preserved in the reproducibility archive.'),
    ('Reference and target inputs',
     "The fixed GSE178341 reference uses layers['counts'], summing repeated gene symbols and retaining71 non-epithelial labels plus Epi. No per-cell total10000 normalization or logarithmic transform is applied to this reference. After the recorded filters,17039 genes remain. Each target cohort is processed separately. Logged RSEM expression is inverted with2**x-1; processed logged microarray expression is inverted with2**x. The continuous target matrices are archived. Before Prism.new, the complete linear target matrix is converted with numpy.rint(...).astype(int64), including nearest-even rounding of half-integers. This integer matrix is used for initial Gibbs inference, cohort-level reference update and final Gibbs inference, rather than merely rounding a temporary sampling total. Microarray intensity proxies are not RNA-seq read counts."),
    ('Sampling and classifier abundance input',
     'The generic mixture branch uses key=None and update_gibbs=True. Chain length is configured as1000, burn-in500, thin2, alpha1 and MAP sigma2. The actual preserved loop updates indices0..998 and retains250 states,500,502,...,998. All explicitly configured run seeds are2026. Final posterior-mean abundance is posterior_theta_f.theta, exported as final_fractions.csv. The parent denominator is the sum of all subtypes belonging to the same clTopLevel in this final matrix. The classifier uses subtype abundance divided by this parent sum. Initial fractions and DECODE outputs are not used in the current twelve-subtype classifier.'),
    ('MMR4 input scale and cutoff ordering',
     'MMR4 takes the four gene values from stored expression.csv matrices before BayesPrism-specific inverse transforms or rounding. TCGA uses log2(RSEM normalized expression+1); the seven GEO cohorts use processed log2 microarray expression. All4848 archived pre-Z-score values across1212 samples were matched to those matrix values. Cohort-specific mean and sample standard deviation (ddof=1) are computed without outcome labels. Cutoff candidates are ranked by descending training macro-AUROC rounded to12 decimals, then descending training macro-BA rounded to12 decimals, then ascending MLH1/MSH2/MSH6/PMS2 cutoff values. Thus the lexicographically smallest tied vector is retained; rounding is part of comparison, not solely display.'),
    ('Single-cell units and composition controls',
     'Reported development AUROC uses all64 specimen predictions directly. Cross-validation groups the62 patient IDs; outcome-stratified bootstrap samples patient clusters and includes every specimen of a selected patient. The main intervals, paired intervals, S7 control intervals and separate centered posthoc p-values use specimen-level AUROC within these cluster resamples. No patient-mean score aggregation is applied to the reported AUROC. CNV sampling draws100000 cells from all cells belonging to QC-passing tumor-tissue specimens, before malignant-candidate assignment. S7 reports fraction-only, subtype-plus-fraction, separate mean/fraction and product-weighted controls. No archived direct add-MSH2/MSH6/PMS2 comparison of the current weighted MSI/dMMR classifier was located; a separate stored-MLH1-loss analysis has a different target and does not supply that evidence.'),
    ('Spatial posterior summaries and expression reconstruction',
     "The reference uses raw counts and73 fixed types. Its gene pool is the intersection of reference symbols, the first spatial section and a fixed historical bulk-common list;16005 unique symbols precede the reference filter (cell_count_cutoff15,cell_percentage_cutoff2=0.03,nonz_mean_cutoff=1.12), which retains12079 model genes. All14 sections contain this same ordered set. Cell2location.export_posterior exports200 posterior draws. Reconstruction inserts post_sample_means for w_sf,m_g and s_g_gene_add into the official conditional expected per-cell-type count formula; reference signature is means_per_cluster_mu_fg. Abundance is means_cell_abundance_w_sf. These are posterior means, not lower quantiles. Reconstructions are computed from parameter means, rather than reconstructing separately for every draw and then averaging."),
    ('Spatial reconstruction formulas',
     'Let X_ig be the observed raw spot count, mu_gf the fixed reference signature, w_if the posterior-mean abundance, m_g the posterior-mean gene scaling, and s_g the posterior-mean additive background. Candidate allocated counts are C_ig = X_ig*(w_i,cand*mu_g,cand*m_g)/(sum_f(w_if*mu_gf)*m_g+s_g). ReconstructedMLH1_i = ln(1+10000*C_i,MLH1/sum_{g in G}C_ig), where G contains all12079 filtered model genes. The section feature is sum_i(ReconstructedMLH1_i*CandidateAbundance_i)/sum_i(CandidateAbundance_i+OtherEpithelialAbundance_i). Thus conditional count allocation precedes candidate-specific CP10K, natural log1p and epithelial abundance weighting. The feature and18 parent-normalized subtype abundances enter the frozen training scaler and shrinkage LDA; two section probabilities are averaged per patient.'),
    ('Spatial training records',
     "The180-epoch reference checkpoint has an empty history_ dictionary and no current-run reference loss trajectory was recovered from retained logs. The is_trained flag does not demonstrate convergence. All14 section checkpoints contain500 loss entries matching their saved CSV records. The relative change between the last50 and preceding50 loss means is0.3184%-1.5452%. Finite histories and loss decreases provide optimization diagnostics but do not establish a prespecified convergence criterion. This limitation is retained. Independently replaying20733 spots from raw10x counts and saved posterior means reproduced reconstructed log expression within5.71e-8, weighted section features within3.07e-9 and frozen probabilities within2.22e-16; this is computational replay, not measurement-equivalence validation."),
    ('Paired fusion-versus-MMR4 AUROC increment',
     'The contrast is AUROC(Fusion05)-AUROC(MMR4_LOCO), using the same1212 samples and their corresponding archived cohort-held-out MMR4 cutoff locks. Fusion05 is0.5*actual twelve-subtype probability+0.5*MMR4_LOCO. Subtype probabilities are five-fold OOF in the two development cohorts and full-development-model predictions in the six GEO cohorts.2000 outcome-stratified within-cohort sample draws use identical indices for both scores. The existing2026 advancing RNG stream was independently reproduced; no seed offsets or loop resets were used. TCGA patient identifiers are distinct; GEO patient clustering has not been independently audited, so the unit is bulk sample. At each replicate, cohort differences are averaged equally within each fixed cohort subset; cohorts themselves are not resampled.95% intervals are percentile2.5/97.5. All upstream measurements, scalers, cutoff locks and classifier predictions remain fixed; no feature selection or fitting uncertainty is included. All32000 bootstrap component AUROCs were independently checked by weighted positive-negative pair comparisons. Point differences and intervals use unrounded values.'),
]
sections = [(title, prose_spacing(content)) for title, content in sections]
W = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'
q = lambda n: '{' + W + '}' + n
document = E.Element(q('document'), nsmap={'w': W}); body = E.SubElement(document, q('body'))
def paragraph(value, bold=False):
    p = E.SubElement(body, q('p')); pp = E.SubElement(p, q('pPr'))
    E.SubElement(pp, q('spacing'), {q('line'): '480', q('lineRule'): 'exact', q('after'): '120', q('before'): '120' if bold else '0'})
    r = E.SubElement(p, q('r')); rp = E.SubElement(r, q('rPr'))
    E.SubElement(rp, q('rFonts'), {q('ascii'): 'Times New Roman', q('hAnsi'): 'Times New Roman'})
    if bold: E.SubElement(rp, q('b'))
    E.SubElement(rp, q('color'), {q('val'): '000000'}); E.SubElement(rp, q('sz'), {q('val'): '24'})
    E.SubElement(r, q('t')).text = value
paragraph('Supplementary methods: implementation and conditional paired comparisons', True)
for title, content in sections: paragraph(title, True); paragraph(content)
sect = E.SubElement(body, q('sectPr'))
E.SubElement(sect, q('pgSz'), {q('w'): '11906', q('h'): '16838'})
E.SubElement(sect, q('pgMar'), {q(k): '1440' for k in ('top', 'right', 'bottom', 'left')})
styles = E.Element(q('styles'), nsmap={'w': W})
normal = E.SubElement(styles, q('style'), {q('type'): 'paragraph', q('default'): '1', q('styleId'): 'Normal'})
E.SubElement(normal, q('name'), {q('val'): 'Normal'})
def xml(node): return E.tostring(node, xml_declaration=True, encoding='UTF-8', standalone=True)
entries = {
    '[Content_Types].xml': b'<?xml version="1.0" encoding="UTF-8"?><Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"><Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/><Default Extension="xml" ContentType="application/xml"/><Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/><Override PartName="/word/styles.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.styles+xml"/></Types>',
    '_rels/.rels': b'<?xml version="1.0" encoding="UTF-8"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/></Relationships>',
    'word/document.xml': xml(document), 'word/styles.xml': xml(styles),
    'word/_rels/document.xml.rels': b'<?xml version="1.0" encoding="UTF-8"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rIdStyles" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/styles" Target="styles.xml"/></Relationships>',
}
with ZipFile(OUT / 'SupplementaryMethods_English.docx', 'w', ZIP_DEFLATED) as z:
    for name, value in entries.items(): z.writestr(name, value)
(OUT / 'SupplementaryMethods_English.txt').write_text('\n\n'.join(title + '\n' + content for title, content in sections))
shutil.copy2(HERE / 'paired_increment/bootstrap_indices.npz', OUT / 'S13_fusion_MMR4_bootstrap_indices.npz')
shutil.copy2(HERE / 'paired_increment/bootstrap_index_metadata.json', OUT / 'S13_fusion_MMR4_bootstrap_metadata.json')
shutil.copy2(HERE / 'paired_increment/bootstrap_AUROC_distributions.npz', OUT / 'S13_fusion_MMR4_bootstrap_AUROCs.npz')
(OUT / 'README_method_update.txt').write_text('Current twelve-subtype classifier, existing cohort-held-out MMR4 and fixed0.5 fusion are unchanged. This update adds implementation provenance, spatial reconstruction details and paired conditional fusion-minus-MMR4 intervals. The specimen-level scRNA AUROC definition and S7 controls are clarified. No model retraining or independent external validation is claimed.\n')
print('Saved English methods and workbook:', len(book.sheetnames), 'sheets')
