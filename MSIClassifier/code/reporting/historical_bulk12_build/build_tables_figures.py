"""English tables and figures for the actually fitted twelve-feature model."""
from pathlib import Path
from zipfile import ZipFile
from copy import deepcopy
import json,hashlib,re,shutil
import numpy as np
import pandas as pd
import openpyxl
from openpyxl.styles import Font,PatternFill,Alignment
from openpyxl.utils import get_column_letter
from lxml import etree as E
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib.colors import to_hex
import cairosvg

HERE=Path(__file__).resolve().parent
ARTICLE=HERE.parent
OLD=ARTICLE/'scRNA_AUROC_threshold_all_candidates_seed2026_20261013'
MODEL=ARTICLE/'bulk_7_vs12_mlp_mmr4_loco_seed2026_20261013'
SUP=HERE/'supplement';FIG=HERE/'figures';HISTORY=SUP/'historical_bulk10'
SHA=lambda path:hashlib.sha256(Path(path).read_bytes()).hexdigest()
COHORTS=['GSE39582','TCGA-COADREAD','GSE24550','GSE25071','GSE27544','GSE35896','GSE39084','GSE75316']
SCOPE={'development_OOF':'Two subtype-development cohorts (OOF)',
       'retrospective':'Six retrospective GEO cohorts','all_eight':'All eight retrospective cohorts',
       'development_apparent':'Two development cohorts (apparent)'}
COMPONENT={'Subtype':'Twelve-subtype MLP','MMR4_LOCO':'Cohort-held-out MMR4','Fusion05':'Fusion 0.50:0.50'}

def save(df,path):df.to_csv(path,index=False,float_format='%.17g')
def dump(path,obj):path.write_text(json.dumps(obj,indent=2,ensure_ascii=False,allow_nan=False)+'\n')
def export(fig,name):
    for ext in ['png','svg','pdf']:fig.savefig(FIG/(name+'.'+ext),dpi=260,bbox_inches='tight')
    plt.close(fig)

def initialize():
    SUP.mkdir(exist_ok=True);FIG.mkdir(exist_ok=True);HISTORY.mkdir(exist_ok=True)
    for f in (OLD/'supplement').iterdir():
        if f.is_file():shutil.copy2(f,SUP/f.name)
    historical_names=['S10.csv','S11.csv','S13.csv','S13_macro_metrics.csv','S13_classifier_ranges.csv','S14.csv','S14_panel_summary.csv',
       'S14_replicate_overlap.csv','S14_selected_count_distribution.csv','S14_threshold_sensitivity.csv','FigureS2_candidate_frequencies.csv',
       'FigureS2_exact_set_recovery.csv','FigureS2_ten_selected_frequency.csv','Parameters.csv','Readme.csv','Software.csv',
       'all_classifier_predictions.csv','primary_MLP_predictions.csv','bootstrap_resampling_records.npz','development_inputs.csv',
       'retrospective_inputs.csv','manuscript_primary_protocol.json','current_bulk_computation_verification.json','SupplementaryTablesS1-S14_English.xlsx']
    for name in historical_names:shutil.copy2(OLD/'supplement'/name,HISTORY/name)
    # These archived10/11-classifier-only records have no current12 identity.
    for name in ['S13_classifier_ranges.csv','S14_threshold_sensitivity.csv','FigureS2_ten_selected_frequency.csv','all_classifier_predictions.csv','prepare_evidence.py','evidence_checks.json']:
        (SUP/name).unlink(missing_ok=True)
    for pattern in ['SupplementaryFigureS1*','SupplementaryFigureS2*']:
        for f in SUP.glob(pattern):f.unlink()
    (HISTORY/'README.txt').write_text('Unchanged historical ten-feature and eleven-classifier records. Not predictions or selection frequencies of the adopted twelve-feature model. Original source: '+str(OLD/'supplement')+'\n')
    package=HERE/'model_computation';package.mkdir(exist_ok=True)
    for f in MODEL.iterdir():
        if f.is_file() and f.suffix not in ['.zip','.log']:shutil.copy2(f,package/f.name)
    for folder in ['inputs','source_archive','models']:
        shutil.copytree(MODEL/folder,package/folder,dirs_exist_ok=True)
    archive=HERE/'source_archive';archive.mkdir(exist_ok=True)
    source_doc=next(OLD.glob('*.docx'));shutil.copy2(source_doc,archive/'source_manuscript.docx')
    dump(HERE/'source_manifest.json',{'manuscript':{'original':str(source_doc),'sha256':SHA(source_doc)},
         'prior_workbook':{'original':str(OLD/'supplement/SupplementaryTablesS1-S14_English.xlsx'),'sha256':SHA(OLD/'supplement/SupplementaryTablesS1-S14_English.xlsx')},
         'model':{'original':str(MODEL),'verification_sha256':SHA(MODEL/'verification.json')},
         'stability':{'protocol_sha256':SHA(HERE/'stability/protocol.json'),'verification_sha256':SHA(HERE/'stability/verification.json')}})

def tables():
    pred=pd.read_csv(MODEL/'predictions.csv');pred=pred[pred.panel=='panel12'].copy()
    metrics=pd.read_csv(MODEL/'cohort_metrics.csv');metrics=metrics[metrics.panel=='panel12'].copy()
    macro=pd.read_csv(MODEL/'macro_metrics.csv');macro=macro[macro.panel=='panel12'].copy()
    measures=['AUROC','BA','sensitivity','specificity','Brier']
    all8=metrics[metrics.split!='development_apparent'].groupby('component')[measures].mean().reset_index()
    all8['panel']='panel12';all8['split']='all_eight';all8['n_cohorts']=8
    macro=pd.concat([macro,all8],ignore_index=True)
    membership=pd.read_csv(MODEL/'panel_membership.csv');panel=json.loads((MODEL/'panels_locked_before_fit.json').read_text())['panel12']
    stability=pd.read_csv(HERE/'stability/S14.csv').set_index('feature')
    summary=json.loads((HERE/'stability/summary.json').read_text())
    table4=[]
    for f in panel:
        row=stability.loc[f]
        table4.append([f,'Positive-associated' if row.scRNA_direction==1 else 'Negative-associated',
                       f'{row.scRNA_oriented_AUROC:.3f}',f'{row["TCGA-COADREAD_oriented_AUROC"]:.3f}',
                       f'{row.GSE39582_oriented_AUROC:.3f}',f'{100*row.bootstrap_inclusion_frequency:.2f}%'])
    table4=pd.DataFrame(table4,columns=['Subtype','Reference direction','Reference oriented AUROC','TCGA oriented AUROC','GSE39582 oriented AUROC','Bootstrap inclusion'])
    table2=[]
    for split in ['development_OOF','retrospective','all_eight']:
        for component in ['Subtype','MMR4_LOCO','Fusion05']:
            row=macro[(macro.split==split)&(macro.component==component)].iloc[0]
            table2.append([COMPONENT[component],SCOPE[split],f'{row.AUROC:.4f}',f'{row.BA:.4f}',f'{row.Brier:.4f}'])
    table2=pd.DataFrame(table2,columns=['Component','Cohort subset','Macro-AUROC','Macro-BA','Macro-Brier'])
    save(table2,HERE/'Table2B.csv');save(table4,HERE/'Table4.csv')
    save(metrics,SUP/'S13.csv');save(macro,SUP/'S13_macro_metrics.csv');save(pred,SUP/'primary_MLP_predictions.csv')
    save(pred[pred.split=='retrospective'],SUP/'retrospective_inputs.csv')
    save(pd.read_csv(MODEL/'inputs/all_cohort_features.csv').query('scope == "development"')[['dataset','sampleID','y','fold']+panel],SUP/'development_inputs.csv')
    save(pd.read_csv(MODEL/'paired_differences.csv'),SUP/'S13_paired_7_vs12.csv')
    for file in ['S14.csv','S14_replicate_overlap.csv','S14_selected_count_distribution.csv']:shutil.copy2(HERE/'stability'/file,SUP/file)
    shutil.copy2(HERE/'stability/resampling_records.npz',SUP/'bootstrap_resampling_records.npz')
    summary_table=pd.DataFrame([{'statistic':key,'value':value,'scope':'17 fixed source candidates; development-only conditional selection sensitivity'}
                               for key,value in summary.items() if isinstance(value,(float,int)) and not isinstance(value,bool)])
    save(summary_table,SUP/'S14_panel_summary.csv')
    save(stability.reset_index(),SUP/'FigureS2_candidate_frequencies.csv')
    save(pd.DataFrame([{'criterion':'exact_panel12','frequency':summary['exact_panel_recovery'],'replicates':summary['n_exact_panel_recovery'],'total':2000},
                       {'criterion':'selected_count_12','frequency':summary['size12_frequency'],'replicates':summary['n_size12'],'total':2000}]),SUP/'FigureS2_exact_set_recovery.csv')
    ledger=pd.read_csv(OLD/'supplement/S10.csv').rename(columns={'retained_bulk10':'historical_retained_bulk10'})
    ledger['candidate_bulk17']=ledger.feature.isin(membership.feature)
    ledger['retained_bulk12']=ledger.feature.isin(panel)
    save(ledger,SUP/'S10.csv')
    raw=pd.read_csv(MODEL/'inputs/all_cohort_features.csv')
    direction=membership.set_index('feature').scRNA_direction
    from sklearn.metrics import roc_auc_score
    single=[]
    for ds,part in raw.groupby('dataset'):
        for f in panel:
            au=roc_auc_score(part.y,part[f]*direction[f])
            single.append({'feature':f,'dataset':ds,'scope':'development' if ds in COHORTS[:2] else 'retrospective',
                           'n':len(part),'MSI':int(part.y.sum()),'frozen_source_direction':int(direction[f]),
                           'oriented_AUROC':au,'source_direction_retained':bool(au>.5)})
    single=pd.DataFrame(single);save(single,SUP/'S11.csv')
    parameters=pd.read_csv(OLD/'supplement/Parameters.csv')
    parameters=parameters[parameters.component!='Actual seed-2026 subtype classifier'].copy()
    protocol=json.loads((MODEL/'protocol.json').read_text())
    parameters=pd.concat([parameters,pd.DataFrame([
        ['Bulk twelve-feature classifier','parameters',json.dumps(protocol['classifier']),'Actual six refitted checkpoints, seed2026'],
        ['Bulk twelve-feature classifier','feature_order',';'.join(panel),'Fixed before fitting; source/development-derived'],
        ['Bulk candidate selection','source_oriented_AUROC','strictly >0.70','17 source candidates, exploratory source-screen definition'],
        ['Bulk candidate selection','both_development_oriented_AUROC','strictly >0.50','Directional concordance; not a significance test'],
        ['Bulk model adoption','timing','After seven/twelve comparative results were viewed','Post hoc exploratory configuration, not prospective prespecification'],
        ['Bulk selection bootstrap','replicates','2000','Archived2026 sample indices, advanced RNG independently replayed'],
        ['Bulk fixed fusion','subtype_MMR4_weights','0.5;0.5','User-fixed, unchanged; archived eight-cohort LOCO MMR4'],
    ],columns=parameters.columns)],ignore_index=True)
    save(parameters,SUP/'Parameters.csv')
    current={**protocol,'panel':panel,'primary_classifier':'mlp_4_l2; actually refitted twelve-feature model',
             'model_adoption':'After seven/twelve comparative results were viewed; exploratory',
             'bootstrap_candidates':17,'bootstrap_threshold':.5,'main_figure':'Twelve retained subtype frequencies only'}
    dump(SUP/'manuscript_primary_protocol.json',current)
    shutil.copy2(MODEL/'verification.json',SUP/'current_bulk_computation_verification.json')
    book=openpyxl.load_workbook(OLD/'supplement/SupplementaryTablesS1-S14_English.xlsx')
    sheets={'S10':ledger,'S11':single,'S13':metrics,'S13_macro':macro,'S13_panel_comparison':pd.read_csv(SUP/'S13_paired_7_vs12.csv'),
            'S13_sample_scores':pred,'S14':stability.reset_index(),'S14_panel_summary':summary_table,
            'S14_overlap':pd.read_csv(SUP/'S14_replicate_overlap.csv'),'S14_size_distribution':pd.read_csv(SUP/'S14_selected_count_distribution.csv'),
            'S14_exact_recovery':pd.read_csv(SUP/'FigureS2_exact_set_recovery.csv'),'Parameters':parameters}
    for name in ['S13_ranges','S14_thresholds']:
        if name in book.sheetnames:del book[name]
    for name,frame in sheets.items():
        position=book.sheetnames.index(name) if name in book.sheetnames else len(book.sheetnames)
        if name in book.sheetnames:del book[name]
        ws=book.create_sheet(name,position);ws.append(frame.columns.tolist())
        for values in frame.replace({np.nan:None}).itertuples(index=False,name=None):ws.append(list(values))
        ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
        for cell in ws[1]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='38566B')
        for col in ws.columns:ws.column_dimensions[get_column_letter(col[0].column)].width=min(44,max(14,len(str(col[0].value))+2))
        for row in ws:
            for cell in row:cell.alignment=Alignment(vertical='top',wrap_text=True)
    descriptions={'S13':'Current twelve-feature MLP, archived LOCO MMR4, fixed fusion: cohort metrics',
                  'S13_macro':'Current twelve-feature components: equal-cohort macro metrics',
                  'S14':'All17 source candidates: development bootstrap inclusion at direction criterion>0.50',
                  'S14_overlap':'2000 resamples: panel size, Jaccard and exact twelve-feature set recovery',
                  'S14_size_distribution':'Selected subtype counts under directional concordance',
                  'S13_panel_comparison':'Paired fixed-score seven/twelve comparison; post hoc and conditional',
                  'S13_sample_scores':'Twelve-feature per-sample OOF, apparent and retrospective predictions',
                  'S14_exact_recovery':'Panel size12 and exact panel12 recovery are distinct summaries',
                  'S10':'71 source candidates; historical scRNA selections and current17/12 bulk flags',
                  'S11':'Frozen-direction single-subtype AUROCs for the current12 panel',
                  'Parameters':'Current twelve-feature classifier, source screening and historical common preprocessing'}
    oldread=list(book['Readme'].values);del book['Readme'];ws=book.create_sheet('Readme',0);ws.append(['sheet','content'])
    for name,desc in oldread[1:]:
        if name in ['S13_ranges','S14_thresholds'] or name in descriptions:continue
        ws.append([name,desc])
    for name,desc in descriptions.items():ws.append([name,desc])
    ws.append(['Historical_archive','Unchanged ten-feature / eleven-classifier records reside in historical_bulk10; not current12 results'])
    book.save(SUP/'SupplementaryTablesS1-S14_English.xlsx')
    save(pd.DataFrame(list(ws.values)[1:],columns=['sheet','content']),SUP/'Readme.csv')
    (SUP/'README.txt').write_text('Current manuscript: twelve-feature MLP plus archived eight-cohort LOCO MMR4, fixed0.5 fusion.\n'
        'SC primary remains fixed18 + weightedMLH1; spatial and SC sensitivity tables remain unchanged.\n'
        'S10 adds candidate_bulk17/retained_bulk12. S11/S13/S14 contain current12-oriented results. Historical bulk10 and11-classifier materials are preserved separately.\n'
        'Bulk>0.50 encodes source-direction concordance, not statistical significance. Source screening>0.70 and full-development selection remain outside classifierCV.\n'
        'The12 configuration was adopted after comparative results were examined; cohorts retain retrospective-exposure status.\n'
        'Classifier results and score uncertainty were previously independently verified. Selection resamples retain and replay the same seed2026 indices; no external labels select subtypes.\n')
    return pred,metrics,macro,stability,summary

def figures(metrics,macro,stability,summary):
    plt.rcParams.update({'font.family':'DejaVu Sans','font.size':10,'svg.fonttype':'none','pdf.fonttype':42,
                         'axes.spines.top':False,'axes.spines.right':False})
    colors=['#2878A5','#B57147','#3E9188']
    retained=stability[stability.retained_bulk12].sort_values('bootstrap_inclusion_frequency',ascending=True)
    selected=metrics[metrics.split!='development_apparent']
    values=selected.pivot(index='dataset',columns='component',values='AUROC').loc[COHORTS,['Subtype','MMR4_LOCO','Fusion05']]
    fig,axes=plt.subplots(2,2,figsize=(13.5,10.5),constrained_layout=True)
    ax=axes[0,0];bars=ax.barh(retained.index,100*retained.bootstrap_inclusion_frequency,color=colors[0])
    ax.bar_label(bars,fmt='%.2f%%',padding=3,fontsize=8);ax.set_xlim(0,114);ax.set_xticks([0,25,50,75,100]);ax.set_xlabel('Bootstrap inclusion frequency (%)')
    ax.set_title('A   Twelve retained subtype frequencies',loc='left',weight='bold')
    ax=axes[0,1];im=ax.imshow(values,vmin=.5,vmax=1,cmap='YlGnBu',aspect='auto')
    ax.set_xticks(range(3),['Twelve-subtype MLP','MMR4 LOCO','Fusion 0.50']);ax.tick_params(axis='x',labelsize=9)
    ax.set_yticks(range(8),[c+(' (OOF)' if c in COHORTS[:2] else '') for c in COHORTS]);ax.tick_params(axis='y',labelsize=9)
    for i in range(8):
        for j in range(3):ax.text(j,i,f'{values.iloc[i,j]:.3f}',ha='center',va='center',color='white' if values.iloc[i,j]>.83 else '#20242A',fontsize=9)
    ax.set_title('B   Per-cohort AUROC',loc='left',weight='bold');fig.colorbar(im,ax=ax,shrink=.9,label='AUROC')
    ax=axes[1,0];x=np.arange(3);width=.36
    for k,(split,label) in enumerate([('development_OOF','Subtype-development cohorts (OOF)'),('retrospective','Retrospective GEO cohorts')]):
        v=macro[macro.split==split].set_index('component').loc[['Subtype','MMR4_LOCO','Fusion05'],'AUROC']
        bars=ax.bar(x+(k-.5)*width,v,width,color=colors[k],label=label);ax.bar_label(bars,fmt='%.3f',padding=3,fontsize=9)
    ax.set_xticks(x,['Twelve-subtype MLP','MMR4 LOCO','Fusion 0.50:0.50']);ax.tick_params(axis='x',labelsize=9);ax.set_ylim(0,1.22)
    ax.set_ylabel('Equal-weight cohort macro-AUROC');ax.set_yticks(np.arange(0,1.01,.2));ax.legend(loc='upper left',fontsize=8,frameon=False)
    ax.set_title('C   Fixed-half fusion and components',loc='left',weight='bold')
    ax=axes[1,1];fusion=selected[selected.component=='Fusion05'].set_index('dataset').loc[COHORTS]
    x=np.arange(8)
    ax.bar(x-.18,fusion.sensitivity,.36,color=colors[0],label='Sensitivity');ax.bar(x+.18,fusion.specificity,.36,color=colors[1],label='Specificity')
    ax.set_xticks(x,COHORTS,rotation=45,ha='right');ax.tick_params(axis='x',labelsize=8);ax.set_ylim(0,1.22)
    ax.set_ylabel('Fusion at score threshold 0.5');ax.set_yticks(np.arange(0,1.01,.2));ax.legend(loc='upper left',fontsize=9,frameon=False)
    ax.set_title('D   Fixed-threshold classification',loc='left',weight='bold')
    export(fig,'Figure4_twelve_subtype_MMR4_LOCO')
    save(retained.reset_index(),FIG/'Figure4A_values.csv');save(selected,FIG/'Figure4BD_values.csv');save(macro[macro.split.isin(['development_OOF','retrospective'])],FIG/'Figure4C_values.csv')
    fig,axes=plt.subplots(2,2,figsize=(13,10),constrained_layout=True)
    allfreq=stability.sort_values('bootstrap_inclusion_frequency',ascending=True)
    ax=axes[0,0];bars=ax.barh(allfreq.index,100*allfreq.bootstrap_inclusion_frequency,color=['#2878A5' if b else '#B9C5CE' for b in allfreq.retained_bulk12])
    ax.bar_label(bars,fmt='%.2f%%',padding=3,fontsize=7.5);ax.set_xlim(0,115);ax.set_xticks([0,25,50,75,100]);ax.set_xlabel('Bootstrap inclusion frequency (%)');ax.set_title('A   All 17 fixed candidates',loc='left',weight='bold')
    overlap=pd.read_csv(HERE/'stability/S14_replicate_overlap.csv');sizes=pd.read_csv(HERE/'stability/S14_selected_count_distribution.csv')
    ax=axes[0,1];ax.bar(sizes.selected_count,sizes.frequency*100,color='#2878A5');ax.set_xlabel('Number of selected subtypes');ax.set_ylabel('Resamples (%)');ax.set_xticks(sizes.selected_count)
    ax.set_title('B   Selected subtype counts',loc='left',weight='bold')
    ax=axes[1,0];ax.hist(overlap.Jaccard_with_panel12,bins=np.linspace(0,1,21),color='#B57147',edgecolor='white')
    ax.axvline(summary['Jaccard_median'],color='#555555',ls='--');ax.set(xlabel='Jaccard overlap with the full-data panel12',ylabel='Resamples');ax.set_title('C   Panel overlap',loc='left',weight='bold')
    ax=axes[1,1];bars=ax.bar(['12 selected subtypes','Exact full-data 12-subtype set'],[100*summary['size12_frequency'],100*summary['exact_panel_recovery']],color=['#2878A5','#B57147'])
    ax.bar_label(bars,fmt='%.2f%%',padding=5,fontsize=10);ax.set_ylim(0,105);ax.set_ylabel('Resamples (%)');ax.tick_params(axis='x',labelsize=8)
    ax.set_title('D   Panel size and exact-set recovery',loc='left',weight='bold')
    export(fig,'SupplementaryFigureS2_bulk12_selection_stability')
    # Selection flow: keep the single-cell18 classifier and the separate bulk17 pool explicit.
    fig,ax=plt.subplots(figsize=(12,7.2));ax.set_xlim(0,12);ax.set_ylim(0,7.2);ax.axis('off')
    def box(x,y,w,h,t,color='#EDF3F6'):
        from matplotlib.patches import FancyBboxPatch
        ax.add_patch(FancyBboxPatch((x,y),w,h,boxstyle='round,pad=0.05',fc=color,ec='#52636E',lw=1))
        ax.text(x+w/2,y+h/2,t,ha='center',va='center',fontsize=10,linespacing=1.5)
    def arrow(a,b):ax.annotate('',xy=b,xytext=a,arrowprops={'arrowstyle':'->','color':'#657780','lw':1.5})
    box(.4,5.8,11.2,1,'GSE178341 reference: 71 unmerged subtype candidates (at least 50 cells)\nFrozen cell definitions and source association directions')
    box(.4,3.7,4.4,1.4,'scRNA primary branch\nHistorical source cutoff >0.65: 24 candidates\nHistorical configuration comparison: fixed18\n18 proportions + weighted malignant MLH1')
    box(6.2,3.7,5.4,1.4,'Bulk candidate branch\nSource oriented AUROC >0.70: 17 candidates\nBoth development oriented AUROCs >0.50\n12 direction-concordant subtypes')
    arrow((3,5.8),(2.6,5.15));arrow((9,5.8),(8.9,5.15))
    box(.4,1.9,4.4,1.1,'Grouped classifier OOF + retrospective scoring\nTraining-only Z-score + shrinkage LDA\nFrozen18 model transferred to spatial proxies')
    box(6.2,1.9,5.4,1.1,'GSE39582 + TCGA-COADREAD\nArchived five-fold OOF / full-fit MLP\nTraining-only Z-score, fixed 4-unit MLP')
    arrow((2.6,3.7),(2.6,3.0));arrow((8.9,3.7),(8.9,3.0))
    box(6.2,.25,5.4,.95,'Eight-cohort LOCO MMR4: fit7, score held-out1\nFixed fusion: 0.5 subtype probability + 0.5 MMR4\nTwo development + six retrospective cohorts','#F4EEE9')
    arrow((8.9,1.9),(8.9,1.25))
    ax.text(.5,.65,'Selection stability: 2,000 development-only resamples\nSource screen and measurements fixed\nPanel adoption includes post hoc configuration comparison',fontsize=9,va='center',color='#53646D')
    export(fig,'SupplementaryFigureS1_current_selection')
    for name in ['SupplementaryFigureS1_current_selection','SupplementaryFigureS2_bulk12_selection_stability']:
        for ext in ['png','svg','pdf']:shutil.copy2(FIG/(name+'.'+ext),SUP/(name+'.'+ext))

def edit_vector_figures(metrics):
    source=HERE/'source_archive/source_manuscript.docx';ns={'s':'http://www.w3.org/2000/svg'}
    with ZipFile(source) as z:
        original1=z.read('word/media/image2.svg');original5=z.read('word/media/image12.svg')
    r=E.fromstring(original1)
    edits={'10 fixed features':'12 subtype features','from 24 candidates':'from 17 candidates','11 classifiers':'MLP (4 units)'}
    found=[]
    for t in r.xpath('//s:text',namespaces=ns):
        value=''.join(t.itertext())
        if value in edits:assert len(t)==0;t.text=edits[value];found.append(value)
    assert set(found)==set(edits)
    svg=E.tostring(r,xml_declaration=True,encoding='UTF-8')
    (FIG/'Figure1_study_design_bulk12.svg').write_bytes(svg)
    cairosvg.svg2png(bytestring=svg,write_to=str(FIG/'Figure1_study_design_bulk12.png'),output_width=2400,output_height=1500)
    cairosvg.svg2pdf(bytestring=svg,write_to=str(FIG/'Figure1_study_design_bulk12.pdf'))
    r=E.fromstring(original5);old=E.fromstring(original5)
    axis=r.xpath('//s:g[@id="axes_3"]',namespaces=ns)[0]
    mesh=axis.xpath('./s:g[@id="QuadMesh_3"]/s:path',namespaces=ns);assert len(mesh)==24
    texts=axis.xpath('.//s:text',namespaces=ns)
    cells=[t for t in texts if t.get('y') is not None and abs(float(t.get('y'))-382.441636)<.01]
    assert len(cells)==8
    ordered=['TCGA-COADREAD','GSE39582']+COHORTS[2:]
    values=metrics[(metrics.component=='Fusion05')&(metrics.split!='development_apparent')].set_index('dataset').loc[ordered,'AUROC']
    cmap=plt.get_cmap('YlGnBu')
    for path,t,value in zip(mesh[:8],cells,values):
        path.set('style','fill: '+to_hex(cmap((float(value)-.65)/.35)))
        t.text=f'{value:.3f}';style=re.sub(r';?fill\s*:\s*[^;]+','',t.get('style','')).rstrip(';')
        t.set('style',style+';fill:'+('#ffffff' if value>.88 else '#20242a'))
    for t in texts:
        value=''.join(t.itertext())
        if value=='(ten-panel; 0.50:0.50 fusion)':t.text='(12 subtypes; 0.50:0.50 fusion)'
        if value=='bulk RNA: fixed-ten LOCO fusion':t.text='bulk RNA: twelve-subtype LOCO fusion'
    # Every unrelated axis and literature cell remains byte-equivalent at XML level.
    for ident in ['axes_1','axes_2_new_axes_1','axes_4','axes_5']:
        a=old.xpath('//s:g[@id="'+ident+'"]',namespaces=ns)[0];b=r.xpath('//s:g[@id="'+ident+'"]',namespaces=ns)[0]
        assert E.tostring(a,method='c14n')==E.tostring(b,method='c14n')
    old_mesh=old.xpath('//s:g[@id="QuadMesh_3"]/s:path',namespaces=ns)
    assert all(E.tostring(a,method='c14n')==E.tostring(b,method='c14n') for a,b in zip(old_mesh[8:],mesh[8:]))
    svg=E.tostring(r,xml_declaration=True,encoding='UTF-8')
    (FIG/'Figure5_comparison_bulk12.svg').write_bytes(svg)
    cairosvg.svg2png(bytestring=svg,write_to=str(FIG/'Figure5_comparison_bulk12.png'),output_width=1834,output_height=1712)
    cairosvg.svg2pdf(bytestring=svg,write_to=str(FIG/'Figure5_comparison_bulk12.pdf'))
    save(values.reset_index(),FIG/'Figure5B_study_values.csv')
    dump(HERE/'vector_figure_verification.json',{'status':'passed','Figure1_label_changes':edits,'Figure1_layout':'preserved',
         'Figure5_other_axes':'unchanged XML','Figure5_literature_cells':'unchanged XML','Figure5_study_values':dict(values)})

def main():
    assert json.loads((HERE/'stability/verification.json').read_text())['status']=='passed'
    assert json.loads((MODEL/'verification.json').read_text())['status']=='passed'
    initialize();pred,metrics,macro,stability,summary=tables()
    figures(metrics,macro,stability,summary);edit_vector_figures(metrics)
    dump(HERE/'tables_figures_complete.json',{'status':'complete','n_current_panel_predictions':len(pred),
        'cohort_metric_rows':len(metrics),'panel12_frequency_min':summary['inclusion_min'],'panel12_frequency_max':summary['inclusion_max']})
    print('Built tables and figures; current panel12 predictions:',len(pred))

if __name__=='__main__':main()
